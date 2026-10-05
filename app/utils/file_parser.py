"""
File parsing utilities for HireMind AI platform.
Extracts text from PDF and DOCX files.
"""

from typing import Optional
from pathlib import Path

from pypdf import PdfReader
from docx import Document

from app.core.exceptions import FileUploadError
from app.utils.text_cleaner import clean_text


class FileParser:
    """Parser for extracting text from uploaded files."""

    @staticmethod
    async def extract_text(file_content: bytes, file_extension: str) -> str:
        """
        Extract text from file content based on extension.

        Args:
            file_content: Raw file bytes
            file_extension: File extension (e.g., '.pdf', '.docx')

        Returns:
            Extracted and cleaned text

        Raises:
            FileUploadError: If parsing fails
        """
        ext = file_extension.lower()

        if ext == '.pdf':
            return await FileParser._extract_pdf_text(file_content)
        elif ext == '.docx':
            return await FileParser._extract_docx_text(file_content)
        else:
            raise FileUploadError(f"Unsupported file type: {ext}")

    @staticmethod
    async def _extract_pdf_text(content: bytes) -> str:
        """
        Extract text from PDF content.

        Args:
            content: PDF file bytes

        Returns:
            Extracted text
        """
        try:
            import io

            # Create a file-like object from bytes
            pdf_file = io.BytesIO(content)
            reader = PdfReader(pdf_file)

            # Extract text from all pages
            text_parts = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)

            full_text = '\n\n'.join(text_parts)
            return clean_text(full_text)

        except Exception as e:
            raise FileUploadError(f"Failed to parse PDF: {str(e)}")

    @staticmethod
    async def _extract_docx_text(content: bytes) -> str:
        """
        Extract text from DOCX content.

        Args:
            content: DOCX file bytes

        Returns:
            Extracted text
        """
        try:
            import io

            # Create a file-like object from bytes
            docx_file = io.BytesIO(content)
            doc = Document(docx_file)

            # Extract text from paragraphs
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]

            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = ' '.join(cell.text for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            full_text = '\n\n'.join(paragraphs)
            return clean_text(full_text)

        except Exception as e:
            raise FileUploadError(f"Failed to parse DOCX: {str(e)}")

    @staticmethod
    def get_file_info(content: bytes, filename: str) -> dict:
        """
        Get basic file information.

        Args:
            content: File bytes
            filename: Original filename

        Returns:
            Dictionary with file info
        """
        return {
            "size_bytes": len(content),
            "size_kb": round(len(content) / 1024, 2),
            "extension": Path(filename).suffix.lower(),
            "filename": filename,
        }
